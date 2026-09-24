"""Solar Casa de Bolsa - Portfolio PRO (Book Trading)
Dashboard de cartera en Gs. y USD, leído directamente de CARTERA_SOLAR_CBSA.xlsx

Uso:
    pip install streamlit pandas plotly openpyxl
    streamlit run reporte_cartera.py      # abre en http://localhost:8501

Dejá el Excel en la misma carpeta que este archivo (o subilo desde la barra lateral).
"""
from io import BytesIO
from pathlib import Path

import pandas as pd
import plotly.express as px
import streamlit as st
from openpyxl import load_workbook

NOMBRE_EXCEL = "CARTERA_SOLAR_CBSA.xlsx"

# Sección del reporte -> nombre de la hoja en el Excel
HOJAS = {"Guaraníes": "CARTERA GS", "Dólares": "CARTERA USD"}
ICONOS = {"Guaraníes": "💰", "Dólares": "💵"}

# Columna con la que se mide la cartera (KPIs y gráficos). Por defecto "Valor Nominal",
# igual que el bloque "CARTERA EN REPO" del Excel. Para usar el costo: "Precio de Compra".
COL_CARTERA = "Valor Nominal"

COL_RENTABILIDAD = "CUANTO SUMA CADA INVERSIÓN A LA RENTABILIDAD"
COLS_NUMERICAS = ["TNA", "TNA Remarcada", "Valor Nominal", "Precio de Compra", COL_RENTABILIDAD]
COLS_DETALLE = [
    "Emisor", "TIPO", "ISIN", "MONEDA", "TNA", "TNA Remarcada", 
    "Valor Nominal", "Precio de Compra", "Riesgo", "Reportado", "Reportador", "Vencimiento"
]


# ----------------------------------------------------------------------------- lectura del Excel
def leer_tabla(ws) -> pd.DataFrame:
    """Tabla de inversiones de la hoja (Tabla1 / Tabla2), sin la fila de totales."""
    tablas = ws.tables.items()  # [(nombre, rango), ...]
    if not tablas:
        raise ValueError(f"La hoja '{ws.title}' no tiene una tabla de Excel definida.")
    _, rango = tablas[0]
    filas = [[celda.value for celda in fila] for fila in ws[rango]]
    df = pd.DataFrame(filas[1:], columns=[str(c).strip() for c in filas[0]])
    df = df[df["Emisor"].notna()].copy()  # descarta la fila de totales
    df["Emisor"] = df["Emisor"].astype(str).str.strip()
    for col in COLS_NUMERICAS:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    reportado = df["Reportado"].astype(str).str.strip().str.upper()
    df["Reportado"] = reportado.map(lambda x: "SI" if x in ("SI", "SÍ") else "NO")
    return df


def leer_sobregiro(ws):
    """Busca la etiqueta 'SOBREGIRO' y toma el monto una fila abajo y una columna a la derecha."""
    for fila in ws.iter_rows():
        for celda in fila:
            if isinstance(celda.value, str) and celda.value.strip().upper() == "SOBREGIRO":
                return ws.cell(celda.row + 1, celda.column + 1).value
    return None


@st.cache_data
def calcular(contenido: bytes) -> dict:
    # data_only=True -> lee los valores calculados que Excel guardó (no las fórmulas)
    wb = load_workbook(BytesIO(contenido), data_only=True)
    resultado = {}
    for moneda, hoja in HOJAS.items():
        ws = wb[hoja]
        df = leer_tabla(ws)
        valores = df[COL_CARTERA]
        if valores.isna().all():
            raise ValueError(
                f"'{COL_CARTERA}' viene vacío en '{hoja}'. Abrí el Excel, guardalo "
                "(Ctrl+S) para que se calculen las fórmulas y volvé a cargarlo."
            )

        # TNA promedio ponderada: la misma que calcula el Excel (suma de peso x TNA remarcada)
        if COL_RENTABILIDAD in df.columns and df[COL_RENTABILIDAD].notna().any():
            tna = df[COL_RENTABILIDAD].sum()
        else:
            tna = (valores * df["TNA Remarcada"]).sum() / valores.sum()

        resultado[moneda] = {
            "df": df,
            "sobregiro": leer_sobregiro(ws),
            "cartera": valores.sum(),
            "reporto": valores[df["Reportado"] == "SI"].sum(),
            "tna": tna,
        }
    return resultado


# ----------------------------------------------------------------------------- presentación
def formatear(valor, moneda: str) -> str:
    if valor is None:
        return "s/d"
    if moneda == "Guaraníes":
        return f"Gs. {valor:,.0f}"
    return f"USD {valor:,.2f}"


def a_lo_ancho(func, *args, **kwargs):
    """Compatible con versiones viejas y nuevas de Streamlit."""
    try:
        return func(*args, use_container_width=True, **kwargs)
    except TypeError:
        return func(*args, width="stretch", **kwargs)


def mostrar_moneda(moneda: str, d: dict):
    df = d["df"]
    st.header(f"{ICONOS[moneda]} {moneda}  ·  {len(df)} instrumentos")

    # KPIs Superiores
    col1, col2, col3, col4 = st.columns(4)
    col1.metric("MONTO DE SOBREGIRO", formatear(d["sobregiro"], moneda))
    col2.metric("MONTO DE LA CARTERA", formatear(d["cartera"], moneda))
    col3.metric("MONTO DE LA CARTERA EN REPORTO", formatear(d["reporto"], moneda))
    col4.metric("TNA PROMEDIO PONDERADA", f"{d['tna'] * 100:.2f}%")

    st.markdown("---")

    c1, c2 = st.columns(2)
    with c1:
        st.subheader("📌 COMPOSICIÓN POR EMISOR")
        por_emisor = df.groupby("Emisor", as_index=False)[COL_CARTERA].sum()
        fig_pie = px.pie(por_emisor, names="Emisor", values=COL_CARTERA, hole=0.4)
        a_lo_ancho(st.plotly_chart, fig_pie, key=f"pie_{moneda}")
    with c2:
        st.subheader("📊 DISTRIBUCIÓN POR RIESGO")
        por_riesgo = df.groupby("Riesgo", as_index=False)[COL_CARTERA].sum()
        fig_bar = px.bar(por_riesgo, x="Riesgo", y=COL_CARTERA, color="Riesgo", text_auto=True)
        a_lo_ancho(st.plotly_chart, fig_bar, key=f"bar_{moneda}")

    st.subheader("📋 DETALLE DE POSICIONES DEL BOOK TRADING")
    a_lo_ancho(st.dataframe, df[[c for c in COLS_DETALLE if c in df.columns]])


def buscar_excel():
    for base in (Path(__file__).parent, Path.cwd()):
        ruta = base / NOMBRE_EXCEL
        if ruta.exists():
            return ruta
    return None


def main():
    st.set_page_config(page_title="Solar Casa de Bolsa - Portfolio PRO", layout="wide")

    st.title("🛡️ Solar Casa de Bolsa - Portfolio PRO (Book Trading)")
    st.markdown("Dashboard interactivo de renta fija con actualización manual por instrumentos al cierre del día.")

    # Controles en barra lateral
    st.sidebar.header("⚙️ Controles de Cartera")
    subido = st.sidebar.file_uploader("Cargar otro Excel (opcional)", type=["xlsx"])
    if st.sidebar.button("🔄 Calcular Todo", type="primary"):
        st.cache_data.clear()
        st.toast("¡Cálculos actualizados con éxito!", icon="🔥")

    if subido is not None:
        contenido = subido.getvalue()
    else:
        ruta = buscar_excel()
        if ruta is None:
            st.error(f"No encuentro {NOMBRE_EXCEL}. Ponelo junto a este script o subilo desde la barra lateral.")
            st.stop()
        contenido = ruta.read_bytes()

    try:
        datos = calcular(contenido)
    except Exception as e:
        st.error(f"No se pudo leer el archivo: {e}")
        st.stop()

    for moneda, d in datos.items():
        mostrar_moneda(moneda, d)
        st.markdown("---")


if __name__ == "__main__":
    main()